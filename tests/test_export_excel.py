import pandas as pd
import pytest
from openpyxl import load_workbook

from export_excel import FORMAT_PRIX, exporter


@pytest.fixture
def rapport(tmp_path):
    df = pd.DataFrame({
        "categorie": ["Poetry", "Travel"],
        "titre": ["C", "A"],
        "prix_gbp": [30.0, 10.5],
        "note": [1, 5],
        "en_stock": [False, True],
        "url": ["https://site/c", "https://site/a"],
    })
    kpi = pd.DataFrame({
        "categorie": ["Poetry", "Travel"],
        "nb_livres": [1, 1],
        "prix_moyen": [30.0, 10.5],
        "prix_min": [30.0, 10.5],
        "prix_max": [30.0, 10.5],
        "note_moyenne": [1.0, 5.0],
    })
    chemin = tmp_path / "rapport.xlsx"
    exporter(df, kpi, chemin)
    return load_workbook(chemin)


def test_le_rapport_contient_les_deux_feuilles(rapport):
    assert rapport.sheetnames == ["Indicateurs", "Catalogue"]


def test_la_feuille_catalogue_contient_toutes_les_lignes(rapport):
    feuille = rapport["Catalogue"]
    assert feuille.max_row == 3  # en-tête + 2 livres
    assert [cellule.value for cellule in feuille[1]] == [
        "categorie", "titre", "prix_gbp", "note", "en_stock", "url"]
    assert feuille["B2"].value == "C"


def test_les_entetes_sont_figes_et_les_prix_formates(rapport):
    feuille = rapport["Catalogue"]
    assert feuille.freeze_panes == "A2"
    assert feuille["C2"].number_format == FORMAT_PRIX
