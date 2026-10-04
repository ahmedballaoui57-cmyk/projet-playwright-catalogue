"""Génération du rapport Excel mis en forme avec OpenPyXL."""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

ENTETE_POLICE = Font(bold=True, color="FFFFFF")
ENTETE_FOND = PatternFill("solid", fgColor="1F3864")
FORMAT_PRIX = '#,##0.00 "£"'
LARGEUR_MAX = 60
TOP_GRAPHIQUE = 10


def exporter(df: pd.DataFrame, kpi: pd.DataFrame, chemin: Path) -> None:
    classeur = Workbook()

    feuille_kpi = classeur.active
    feuille_kpi.title = "Indicateurs"
    _ecrire_feuille(feuille_kpi, kpi, colonnes_prix=["prix_moyen", "prix_min", "prix_max"])
    _ajouter_graphique(feuille_kpi, nb_lignes=min(TOP_GRAPHIQUE, len(kpi)))

    feuille_donnees = classeur.create_sheet("Catalogue")
    _ecrire_feuille(feuille_donnees, df, colonnes_prix=["prix_gbp"])

    classeur.save(chemin)


def _ecrire_feuille(feuille: Worksheet, df: pd.DataFrame, colonnes_prix: list[str]) -> None:
    for ligne in dataframe_to_rows(df, index=False, header=True):
        feuille.append(ligne)

    for cellule in feuille[1]:
        cellule.font = ENTETE_POLICE
        cellule.fill = ENTETE_FOND
        cellule.alignment = Alignment(horizontal="center")
    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = feuille.dimensions

    for position, colonne in enumerate(df.columns, start=1):
        lettre = get_column_letter(position)
        plus_long = max(len(str(colonne)), int(df[colonne].astype(str).str.len().max()))
        feuille.column_dimensions[lettre].width = min(plus_long + 2, LARGEUR_MAX)
        if colonne in colonnes_prix:
            for cellule in feuille[lettre][1:]:
                cellule.number_format = FORMAT_PRIX


def _ajouter_graphique(feuille: Worksheet, nb_lignes: int) -> None:
    """Histogramme du nombre de livres pour les catégories les plus fournies."""
    graphique = BarChart()
    graphique.title = f"Top {nb_lignes} des catégories par nombre de livres"
    graphique.y_axis.title = "Nombre de livres"
    graphique.height = 9
    graphique.width = 20

    # Colonne A : catégorie, colonne B : nb_livres (ligne 1 = en-tête).
    valeurs = Reference(feuille, min_col=2, min_row=1, max_row=nb_lignes + 1)
    categories = Reference(feuille, min_col=1, min_row=2, max_row=nb_lignes + 1)
    graphique.add_data(valeurs, titles_from_data=True)
    graphique.set_categories(categories)
    feuille.add_chart(graphique, "H2")
